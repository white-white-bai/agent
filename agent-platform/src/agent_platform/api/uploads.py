"""上传文件的统一读取：大小上限、扩展名白名单、按格式抽取文本（PDF/DOCX/XLSX/XLS/HTML 等）。"""
from __future__ import annotations

import io
import re
from html.parser import HTMLParser
from pathlib import PurePath
from typing import ClassVar

from fastapi import HTTPException, UploadFile

from agent_platform.core.config import get_settings

READ_CHUNK = 64 * 1024

# 需要按格式解析（非纯文本）的扩展名
BINARY_SUFFIXES = {".pdf", ".docx", ".xlsx", ".xls"}


async def read_upload(file: UploadFile, allowed_suffixes: set[str]) -> str:
    """图 DSL 等纯文本上传：只做大小校验与文本解码。"""
    suffix = _check_suffix(file, allowed_suffixes)
    raw = await _read_all(file, get_settings().upload_max_bytes)
    if suffix in BINARY_SUFFIXES:
        raise HTTPException(status_code=400, detail=f"{suffix} 为二进制格式，不能作为 DSL 导入")
    return decode_text(raw)


async def read_document(file: UploadFile, allowed_suffixes: set[str]) -> str:
    """知识库文档上传：大小上限 15MB（可配），按扩展名抽取正文。"""
    suffix = _check_suffix(file, allowed_suffixes)
    raw = await _read_all(file, get_settings().document_max_bytes)
    return extract_text(suffix, raw)


def _check_suffix(file: UploadFile, allowed_suffixes: set[str]) -> str:
    suffix = PurePath(file.filename or "").suffix.lower()
    if suffix not in allowed_suffixes:
        raise HTTPException(status_code=400, detail=f"不支持的文件类型：{suffix or '无扩展名'}")
    return suffix


async def _read_all(file: UploadFile, limit: int) -> bytes:
    buffer = bytearray()
    while chunk := await file.read(READ_CHUNK):
        buffer.extend(chunk)
        if len(buffer) > limit:
            raise HTTPException(status_code=413, detail=f"文件超过上限 {limit} 字节")
    return bytes(buffer)


def extract_text(suffix: str, raw: bytes) -> str:
    try:
        if suffix in (".html", ".htm"):
            return strip_html(decode_text(raw))
        if suffix == ".pdf":
            return _extract_pdf(raw)
        if suffix == ".docx":
            return _extract_docx(raw)
        if suffix == ".xlsx":
            return _extract_xlsx(raw)
        if suffix == ".xls":
            return _extract_xls(raw)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"{suffix} 解析失败：{exc}") from exc
    return decode_text(raw)


def decode_text(raw: bytes) -> str:
    if raw.startswith(b"\xef\xbb\xbf"):
        raw = raw[3:]
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        pass
    try:
        return raw.decode("gb18030")
    except UnicodeDecodeError as exc:
        raise HTTPException(status_code=400, detail="文件编码无法识别，请另存为 UTF-8") from exc


class _HtmlTextCollector(HTMLParser):
    _SKIP: ClassVar[set[str]] = {"script", "style", "head", "noscript", "title"}
    _BREAK: ClassVar[set[str]] = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag in self._SKIP:
            self._skip += 1
        elif tag in self._BREAK:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in self._SKIP and self._skip:
            self._skip -= 1
        elif tag in self._BREAK:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._skip and data.strip():
            self.parts.append(data.strip())


def strip_html(text: str) -> str:
    collector = _HtmlTextCollector()
    collector.feed(text)
    joined = " ".join(collector.parts)
    joined = re.sub(r"[ \t]+", " ", joined)
    return re.sub(r"\n\s*\n+", "\n\n", joined).strip()


def _extract_pdf(raw: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw))
    pages = [(page.extract_text() or "").strip() for page in reader.pages]
    return "\n\n".join(page for page in pages if page)


def _extract_docx(raw: bytes) -> str:
    from docx import Document as DocxDocument

    document = DocxDocument(io.BytesIO(raw))
    lines = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _extract_xlsx(raw: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in workbook.worksheets:
        lines.append(f"# {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if value is None else str(value) for value in row]
            if any(cell.strip() for cell in cells):
                lines.append(" | ".join(cells))
    workbook.close()
    return "\n".join(lines)


def _extract_xls(raw: bytes) -> str:
    import xlrd

    book = xlrd.open_workbook(file_contents=raw)
    lines: list[str] = []
    for sheet in book.sheets():
        lines.append(f"# {sheet.name}")
        for index in range(sheet.nrows):
            cells = [str(value) for value in sheet.row_values(index)]
            if any(cell.strip() for cell in cells):
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def chunk_text(text: str, size: int, overlap: int = 0) -> list[str]:
    """按空行切段、超长硬切、可选重叠。分段实现已迁至 knowledge/chunking.split_text。"""
    from agent_platform.knowledge.chunking import split_text

    return split_text(text, separator="\n\n", max_chars=size, overlap_chars=overlap)
