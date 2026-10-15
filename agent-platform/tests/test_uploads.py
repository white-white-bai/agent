"""上传解析：HTML 去标签、XLSX/DOCX 抽取、分块重叠与坏文件报错。"""
from __future__ import annotations

import io

import pytest
from docx import Document as DocxDocument
from fastapi import HTTPException
from openpyxl import Workbook

from agent_platform.api.uploads import chunk_text, extract_text, strip_html


def test_strip_html_keeps_visible_text_only():
    html = (
        "<html><head><style>p{color:red}</style></head><body>"
        "<h1>标题</h1><p>第一段</p><script>var a = 1</script><p>第二段</p>"
        "</body></html>"
    )
    text = strip_html(html)
    assert "标题" in text and "第一段" in text and "第二段" in text
    assert "color:red" not in text
    assert "var a" not in text


def test_extract_xlsx_rows_become_lines():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "数据"
    sheet.append(["地区", "数量"])
    sheet.append(["常州", 3])
    buffer = io.BytesIO()
    workbook.save(buffer)

    text = extract_text(".xlsx", buffer.getvalue())
    assert "# 数据" in text
    assert "地区 | 数量" in text
    assert "常州 | 3" in text


def test_extract_docx_paragraphs_and_tables():
    document = DocxDocument()
    document.add_paragraph("第一段")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "甲"
    table.cell(0, 1).text = "乙"
    buffer = io.BytesIO()
    document.save(buffer)

    text = extract_text(".docx", buffer.getvalue())
    assert "第一段" in text
    assert "甲 | 乙" in text


def test_extract_text_decodes_plain_and_html():
    assert extract_text(".md", "标题".encode()) == "标题"
    assert extract_text(".htm", "<p>结束</p>".encode()) == "结束"


def test_malformed_binary_reports_clear_error():
    with pytest.raises(HTTPException) as exc:
        extract_text(".pdf", b"definitely not a pdf")
    assert exc.value.status_code == 400
    assert "解析失败" in exc.value.detail


def test_chunk_text_overlap_prepends_previous_tail():
    text = "A" * 10 + "\n\n" + "B" * 10
    assert chunk_text(text, 12) == ["AAAAAAAAAA", "BBBBBBBBBB"]
    overlapped = chunk_text(text, 12, overlap=4)
    assert overlapped[0] == "AAAAAAAAAA"
    assert overlapped[1] == "AAAA\n\nBBBBBBBBBB"
